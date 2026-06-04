"""
Document Loading and Processing
Handles loading, parsing, and chunking of various document formats
"""

import os
import json
import re
from typing import List, Dict, Tuple, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)

class Document:
    """Represents a single document with metadata"""
    
    def __init__(self, content: str, source: str, doc_type: str, metadata: dict = None):
        self.content = content
        self.source = source
        self.doc_type = doc_type
        self.metadata = metadata or {}
        self.chunks = []
    
    def __repr__(self):
        return f"Document(source={self.source}, type={self.doc_type}, size={len(self.content)})"


class DocumentProcessor:
    """Process various document formats into text"""
    
    @staticmethod
    def load_text(file_path: str) -> str:
        """Load plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error loading text file {file_path}: {e}")
            raise
    
    @staticmethod
    def load_json(file_path: str) -> str:
        """Load JSON file and convert to formatted text"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return json.dumps(data, indent=2)
        except Exception as e:
            logger.error(f"Error loading JSON file {file_path}: {e}")
            raise
    
    @staticmethod
    def load_csv(file_path: str) -> str:
        """Load CSV file"""
        try:
            import csv
            rows = []
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(' | '.join(row))
            return '\n'.join(rows)
        except Exception as e:
            logger.error(f"Error loading CSV file {file_path}: {e}")
            raise
    
    @staticmethod
    def load_pdf(file_path: str) -> str:
        """Load PDF file"""
        try:
            import PyPDF2
            text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text.append(page.extract_text())
            return '\n'.join(text)
        except Exception as e:
            logger.error(f"Error loading PDF file {file_path}: {e}")
            raise
    
    @staticmethod
    def load_docx(file_path: str) -> str:
        """Load DOCX file"""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text.append(para.text)
            return '\n'.join(text)
        except Exception as e:
            logger.error(f"Error loading DOCX file {file_path}: {e}")
            raise
    
    @classmethod
    def load_document(cls, file_path: str, doc_type: str = None) -> Document:
        """
        Load document from file based on extension
        
        Args:
            file_path: Path to the document
            doc_type: Override document type (resume, job_description, transcript, etc.)
        
        Returns:
            Document object with content and metadata
        """
        file_ext = Path(file_path).suffix.lower()[1:]  # Remove the dot
        
        # Determine document type
        if doc_type is None:
            filename = Path(file_path).stem.lower()
            if 'resume' in filename or 'cv' in filename:
                doc_type = 'resume'
            elif 'job' in filename or 'jd' in filename:
                doc_type = 'job_description'
            elif 'transcript' in filename or 'interview' in filename:
                doc_type = 'transcript'
            else:
                doc_type = 'unknown'
        
        # Load based on extension
        if file_ext == 'txt':
            content = cls.load_text(file_path)
        elif file_ext == 'json':
            content = cls.load_json(file_path)
        elif file_ext == 'csv':
            content = cls.load_csv(file_path)
        elif file_ext == 'pdf':
            content = cls.load_pdf(file_path)
        elif file_ext == 'docx':
            content = cls.load_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")
        
        # Create document with metadata
        metadata = {
            'file_name': Path(file_path).name,
            'file_size': os.path.getsize(file_path),
            'file_type': file_ext
        }
        
        return Document(content, file_path, doc_type, metadata)


class TextChunker:
    """Split documents into chunks for embedding"""
    
    @staticmethod
    def split_sentences(text: str) -> List[str]:
        """Split text into sentences"""
        # Simple sentence splitter
        sentences = re.split(r'(?<=[.!?])\s+', text)
        return [s.strip() for s in sentences if s.strip()]
    
    @staticmethod
    def chunk_by_size(text: str, chunk_size: int = 512, overlap: int = 100) -> List[str]:
        """
        Split text into chunks of approximate size with overlap
        
        Args:
            text: Text to chunk
            chunk_size: Target chunk size in characters
            overlap: Overlap between chunks in characters
        
        Returns:
            List of text chunks
        """
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + chunk_size
            
            if end >= len(text):
                chunks.append(text[start:].strip())
                break
            
            # Try to find a good break point (end of sentence)
            break_point = text.rfind('.', start, end)
            if break_point == -1 or break_point < start + chunk_size // 2:
                break_point = text.rfind(' ', start, end)
            if break_point == -1:
                break_point = end
            
            chunks.append(text[start:break_point].strip())
            start = max(break_point - overlap, start + 1)
        
        return [c for c in chunks if c]
    
    @staticmethod
    def chunk_document(document: Document, chunk_size: int = 512, overlap: int = 100) -> List[Dict]:
        """
        Chunk a document and return chunks with metadata
        
        Args:
            document: Document object to chunk
            chunk_size: Target chunk size
            overlap: Overlap between chunks
        
        Returns:
            List of chunk dictionaries with content and metadata
        """
        chunks = TextChunker.chunk_by_size(document.content, chunk_size, overlap)
        
        chunk_objects = []
        for i, chunk in enumerate(chunks):
            chunk_obj = {
                'content': chunk,
                'source': document.source,
                'doc_type': document.doc_type,
                'chunk_id': i,
                'metadata': {
                    **document.metadata,
                    'chunk_index': i,
                    'total_chunks': len(chunks)
                }
            }
            chunk_objects.append(chunk_obj)
        
        return chunk_objects


class DocumentManager:
    """Manage document loading, processing, and caching"""
    
    def __init__(self, upload_dir: str = "./uploads", cache_dir: str = "./cache"):
        self.upload_dir = Path(upload_dir)
        self.cache_dir = Path(cache_dir)
        
        # Create directories if they don't exist
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.cache_dir.mkdir(parents=True, exist_ok=True)
    
    def process_file(self, file_path: str, doc_type: str = None, 
                    chunk_size: int = 512, overlap: int = 100) -> Tuple[Document, List[Dict]]:
        """
        Load, process, and chunk a document
        
        Returns:
            (Document object, List of chunks)
        """
        # Load document
        document = DocumentProcessor.load_document(file_path, doc_type)
        logger.info(f"Loaded document: {document}")
        
        # Chunk document
        chunks = TextChunker.chunk_document(document, chunk_size, overlap)
        logger.info(f"Created {len(chunks)} chunks from document")
        
        return document, chunks
    
    def process_batch(self, file_paths: List[str], chunk_size: int = 512, 
                     overlap: int = 100) -> Dict[str, Tuple[Document, List[Dict]]]:
        """Process multiple documents"""
        results = {}
        
        for file_path in file_paths:
            try:
                doc, chunks = self.process_file(file_path, chunk_size=chunk_size, overlap=overlap)
                results[file_path] = (doc, chunks)
            except Exception as e:
                logger.error(f"Failed to process {file_path}: {e}")
                results[file_path] = (None, None)
        
        return results
